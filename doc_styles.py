"""Utilidades para crear y dar formato al documento Word."""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT


def create_doc():
    doc = Document()
    style = doc.styles["Normal"]
    font = style.font
    font.name = "Calibri"
    font.size = Pt(11)
    font.color.rgb = RGBColor(0, 0, 0)
    return doc


def add_title(doc, text):
    p = doc.add_heading(text, level=0)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p


def add_h1(doc, text):
    return doc.add_heading(text, level=1)


def add_h2(doc, text):
    return doc.add_heading(text, level=2)


def add_paragraph(doc, text, bold=False):
    p = doc.add_paragraph(text)
    if bold:
        for run in p.runs:
            run.bold = True
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(text, style="List Bullet")
    if level > 0:
        p.paragraph_format.left_indent = Cm(1.5 * level)
    return p


def add_table_2col(doc, rows):
    """rows = list of (label, value) tuples."""
    table = doc.add_table(rows=len(rows), cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (label, value) in enumerate(rows):
        c0 = table.cell(i, 0)
        c1 = table.cell(i, 1)
        c0.text = str(label)
        c1.text = str(value)
        for p in c0.paragraphs:
            for r in p.runs:
                r.bold = True
    return table


def add_table_3col(doc, rows):
    """rows = list of (col1, col2, col3) tuples."""
    table = doc.add_table(rows=len(rows), cols=3)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row_data in enumerate(rows):
        for j, val in enumerate(row_data):
            table.cell(i, j).text = str(val)
    return table


def add_table_4col(doc, rows):
    """rows = list of (c1, c2, c3, c4) tuples."""
    table = doc.add_table(rows=len(rows), cols=4)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row_data in enumerate(rows):
        for j, val in enumerate(row_data):
            table.cell(i, j).text = str(val)
    return table


def add_spacer(doc):
    doc.add_paragraph("")
