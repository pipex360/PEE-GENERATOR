"""
Generador de PEE basado en plantilla.
Copia el documento plantilla y reemplaza los placeholders con los datos del formulario.
"""
import os
import shutil
import re
from docx import Document


TEMPLATE_PATH = os.path.join(os.path.dirname(__file__), "plantilla_pee.docx")

# Factores de carga de ocupación según Art. 4.2.4 OGUC (m² por persona)
CARGA_FACTORES = {
    "Habitacional": 15,
    "Oficina": 10,
    "Comercial": 3,
    "Mixto": 10,
}


def _calcular_carga_ocupacion(data):
    """Calcula la carga de ocupación según Art. 4.2.4 OGUC.
    Fórmula: superficie / factor según destino."""
    superficie_str = data.get("superficie", "")
    destino = data.get("destino_edificacion", "Habitacional")

    # Limpiar superficie: "2.670,55" -> 2670.55
    superficie_str = superficie_str.replace(".", "").replace(",", ".")
    try:
        superficie = float(superficie_str)
    except (ValueError, TypeError):
        return ""

    factor = CARGA_FACTORES.get(destino, 10)
    carga = int(superficie / factor)
    return str(carga)


def _calcular_destino_pisos_sup(data):
    """Calcula el rango de pisos superiores: 'Piso 2 al {N}'."""
    pisos_str = data.get("pisos_sobre", "")
    # Extraer número de la cadena (ej: "14 pisos" -> 14, "14" -> 14)
    match = re.search(r"(\d+)", pisos_str)
    if match:
        n = int(match.group(1))
        if n > 2:
            return f"Piso 2 al {n}"
        elif n == 2:
            return "Piso 2"
    return data.get("destino_pisos_sup", "")


def _build_replacements(data):
    """Construye el diccionario de reemplazos placeholder -> valor."""
    direccion = data.get("direccion", "")
    acceso_carro = data.get("acceso_carro_bomba", "NO")

    # Auto-calcular carga de ocupación
    carga = _calcular_carga_ocupacion(data)

    # Auto-calcular rango pisos superiores
    pisos_sup = _calcular_destino_pisos_sup(data)

    return {
        # Portada
        "{{REALIZADO_POR}}": data.get("realizado_por", ""),
        "{{APROBADO_POR}}": data.get("aprobado_por", ""),
        "{{FECHA_ACTUALIZACION}}": data.get("fecha_actualizacion", ""),
        # Identificación
        "{{DIRECCION_UPPER}}": direccion.upper(),
        "{{DIRECCION}}": direccion,
        "{{COORDENADAS}}": data.get("coordenadas", ""),
        "{{ENTRE_CALLE_1}}": data.get("entre_calle_1", ""),
        "{{ENTRE_CALLE_2}}": data.get("entre_calle_2", ""),
        "{{ACCESO_EDIFICIO}}": data.get("acceso_edificio", ""),
        "{{PERMISO_MUNICIPAL}}": data.get("permiso_municipal", "Pendiente"),
        "{{ROL_AVALUOS}}": data.get("rol_avaluos", "Pendiente"),
        "{{COMUNA}}": data.get("comuna", ""),
        # Características
        "{{PISOS_SOBRE}}": data.get("pisos_sobre", ""),
        "{{PISOS_BAJO}}": data.get("pisos_bajo", ""),
        "{{SUPERFICIE}}": data.get("superficie", ""),
        "{{CARGA_OCUPACION}}": carga,
        # Carro bomba: SI X / NO o SI / NO X
        "{{CARRO_SI}}": "SI X" if acceso_carro == "SI" else "SI",
        "{{CARRO_NO}}": "NO" if acceso_carro == "SI" else "NO X",
        "{{CALLE_CARRO_BOMBA}}": data.get("calle_carro_bomba", "") if acceso_carro == "SI" else "",
        "{{APERTURAS_EXTERIOR}}": data.get("aperturas_exterior", "Móviles"),
        "{{NUM_UNIDADES}}": data.get("num_unidades", ""),
        "{{NUM_ESTACIONAMIENTOS}}": data.get("num_estacionamientos", "No"),
        "{{DESTINO_EDIFICACION}}": data.get("destino_edificacion", ""),
        # Destinos por piso
        "{{DESTINO_PISO_SUB}}": data.get("destino_piso_sub", "Subterráneo"),
        "{{DESTINO_PISO_SUB_DESC}}": data.get("destino_piso_sub_desc", ""),
        "{{DESTINO_PISO_1}}": data.get("destino_piso_1", "Piso 1"),
        "{{DESTINO_PISO_1_DESC}}": data.get("destino_piso_1_desc", ""),
        "{{DESTINO_PISOS_SUP}}": pisos_sup,
        "{{DESTINO_PISOS_SUP_DESC}}": data.get("destino_pisos_sup_desc", ""),
        # Estructura
        "{{CLASE_ESTRUCTURA}}": data.get("clase_estructura", ""),
        "{{DESCRIPCION_ESTRUCTURA}}": data.get("descripcion_estructura", ""),
        "{{TABIQUES_INTERIORES}}": data.get("tabiques_interiores", ""),
        "{{FACHADAS_EXTERIORES}}": data.get("fachadas_exteriores", ""),
        # Alarmas
        "{{BOCINAS_ALARMA}}": data.get("bocinas_alarma", "No"),
        "{{DETECTORES_HUMO}}": data.get("detectores_humo", "No"),
        "{{DETECTORES_CALOR}}": data.get("detectores_calor", "No aplica"),
        "{{PALANCAS_ALARMA}}": data.get("palancas_alarma", "No"),
        # Comunicación
        "{{TELEFONOS}}": data.get("telefonos", ""),
        "{{CITOFONOS}}": data.get("citofonos", ""),
        "{{ALTAVOCES}}": data.get("altavoces", "No"),
        "{{ALTAVOCES_EVAC}}": data.get("altavoces", "No"),
        "{{OTROS_COMUNICACION}}": data.get("otros_comunicacion", "No"),
        # Incendios
        "{{RED_SECA}}": data.get("red_seca", "No"),
        "{{RED_HUMEDA}}": data.get("red_humeda", ""),
        "{{ESTANQUE_AGUA}}": data.get("estanque_agua", "No"),
        "{{EXTINTORES}}": data.get("extintores", "No"),
        "{{RED_INERTE}}": data.get("red_inerte", "No"),
        # Evacuación
        "{{VIAS_EVACUACION}}": data.get("vias_evacuacion", ""),
        "{{PUNTO_REUNION}}": data.get("punto_reunion", ""),
        "{{ZONA_SEGURIDAD}}": data.get("zona_seguridad", ""),
        "{{ZONA_SEGURIDAD_TEXTO}}": data.get("zona_seguridad", ""),
        # Electricidad
        "{{TABLERO_GENERAL}}": data.get("tablero_general", "SI"),
        "{{TABLEROS_UNIDADES}}": data.get("tableros_unidades", ""),
        "{{GRUPO_ELECTROGENO}}": data.get("grupo_electrogeno", "No"),
        "{{ILUMINACION_EMERGENCIA}}": data.get("iluminacion_emergencia", ""),
        # Gas
        "{{GAS}}": data.get("gas", "No"),
        "{{MEDIDORES_GAS}}": data.get("medidores_gas", "No"),
        "{{TANQUE_GAS}}": data.get("tanque_gas", "No aplica"),
        "{{TANQUE_PETROLEO}}": data.get("tanque_petroleo", "No"),
        # Ventilación
        "{{VENTILACION_CENTRALIZADA}}": data.get("ventilacion_centralizada", "No"),
        "{{TABLERO_COMANDO_VENT}}": data.get("tablero_comando_vent", "No"),
        "{{TOMA_AIRE}}": data.get("toma_aire", "No"),
        # Ascensores
        "{{NUM_ASCENSORES}}": data.get("num_ascensores", ""),
        "{{CAPACIDAD_PERSONAS}}": data.get("capacidad_personas", ""),
        "{{CAPACIDAD_KILOS}}": data.get("capacidad_kilos", ""),
        "{{SISTEMA_ASCENSOR}}": data.get("sistema_ascensor", "Eléctrico"),
        "{{LLAVE_BOMBEROS}}": data.get("llave_bomberos", "Conserjería"),
        # Teléfonos emergencia
        "{{CESFAM_NOMBRE}}": data.get("cesfam_nombre", ""),
        "{{CESFAM_TELEFONO}}": data.get("cesfam_telefono", ""),
        "{{SAPU_NOMBRE}}": data.get("sapu_nombre", ""),
        "{{SAPU_TELEFONO}}": data.get("sapu_telefono", ""),
        "{{HOSPITAL_NOMBRE}}": data.get("hospital_nombre", ""),
        "{{HOSPITAL_TELEFONO}}": data.get("hospital_telefono", ""),
        "{{HOSPITAL_URGENCIA_NOMBRE}}": data.get("hospital_urgencia_nombre", ""),
        "{{HOSPITAL_URGENCIA_TELEFONO}}": data.get("hospital_urgencia_telefono", ""),
        "{{BOMBEROS_CUARTEL}}": data.get("bomberos_cuartel", ""),
        "{{BOMBEROS_TELEFONO}}": data.get("bomberos_telefono", ""),
        "{{COMISARIA_NOMBRE}}": data.get("comisaria_nombre", ""),
        "{{COMISARIA_TELEFONO}}": data.get("comisaria_telefono", ""),
        "{{PLAN_CUADRANTE}}": data.get("plan_cuadrante", ""),
        "{{SEGURIDAD_CIUDADANA}}": data.get("seguridad_ciudadana", ""),
        # Servicios edificio
        "{{EMPRESA_BOMBAS_AGUA}}": data.get("empresa_bombas_agua", ""),
        "{{EMPRESA_ELECTROGENO}}": data.get("empresa_electrogeno", ""),
        "{{EMPRESA_ASCENSORES}}": data.get("empresa_ascensores", ""),
        "{{EMPRESA_INCENDIOS}}": data.get("empresa_incendios", ""),
        "{{EMPRESA_EXTINTORES}}": data.get("empresa_extintores", ""),
    }


def _replace_in_runs(runs, replacements):
    """Reemplaza placeholders en los runs de un párrafo."""
    for run in runs:
        for placeholder, value in replacements.items():
            if placeholder in run.text:
                run.text = run.text.replace(placeholder, value)


def generate_pee(data, output_path):
    """Genera el documento PEE copiando la plantilla y reemplazando los datos."""
    shutil.copy2(TEMPLATE_PATH, output_path)
    doc = Document(output_path)
    replacements = _build_replacements(data)

    # Reemplazar en párrafos
    for paragraph in doc.paragraphs:
        _replace_in_runs(paragraph.runs, replacements)

    # Reemplazar en tablas
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    _replace_in_runs(paragraph.runs, replacements)

    # Reemplazar en headers y footers
    for section in doc.sections:
        for part in [section.header, section.first_page_header,
                     section.footer, section.first_page_footer]:
            if part:
                for paragraph in part.paragraphs:
                    _replace_in_runs(paragraph.runs, replacements)

    doc.save(output_path)
    return output_path
